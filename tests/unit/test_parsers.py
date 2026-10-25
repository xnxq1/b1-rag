from pathlib import Path

import docx
import pymupdf
from docx.enum.text import WD_BREAK

from rag.adapters.outbound.parsing.dispatcher import FileDocumentParser


def _page_of(parsed, offset: int) -> int:
    return next(p.page for p in parsed.pages if p.char_start <= offset <= p.char_end)


async def test_pdf_pages_and_headings(tmp_path: Path) -> None:
    path = tmp_path / "contract.pdf"
    pdf = pymupdf.open()
    fontfile = None
    for candidate in (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if Path(candidate).exists():
            fontfile = candidate
    for text in ("1. ПРЕДМЕТ ДОГОВОРА\nПоставщик передает товар.", "4.2 Штрафные санкции\nНеустойка 0,1% в день."):
        page = pdf.new_page()
        if fontfile:
            page.insert_font(fontname="f", fontfile=fontfile)
            page.insert_text((72, 72), text, fontname="f", fontsize=11)
        else:
            page.insert_text((72, 72), text, fontsize=11)
    pdf.save(path)

    parsed = await FileDocumentParser().parse(path)

    assert parsed.pages_total == 2
    assert [p.page for p in parsed.pages] == [1, 2]
    if fontfile:
        offset = parsed.text.index("Неустойка")
        assert _page_of(parsed, offset) == 2
        assert [h.title for h in parsed.headings] == ["1. ПРЕДМЕТ ДОГОВОРА", "4.2 Штрафные санкции"]


async def test_docx_page_breaks_tables_and_heading_styles(tmp_path: Path) -> None:
    path = tmp_path / "contract.docx"
    document = docx.Document()
    document.add_heading("Предмет договора", level=1)
    document.add_paragraph("Поставщик передает товар.")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("4.2 Штрафные санкции")
    document.add_paragraph("Неустойка 0,1% в день.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Сумма"
    table.rows[0].cells[1].text = "1 500 000 руб."
    document.save(path)

    parsed = await FileDocumentParser().parse(path)

    assert parsed.pages_total == 2
    assert _page_of(parsed, parsed.text.index("Поставщик")) == 1
    assert _page_of(parsed, parsed.text.index("Неустойка")) == 2
    assert [h.title for h in parsed.headings] == ["Предмет договора", "4.2 Штрафные санкции"]
    assert "Сумма | 1 500 000 руб." in parsed.text


async def test_docx_without_breaks_has_no_pages(tmp_path: Path) -> None:
    path = tmp_path / "plain.docx"
    document = docx.Document()
    document.add_paragraph("Текст без разрывов страниц.")
    document.save(path)

    parsed = await FileDocumentParser().parse(path)

    assert parsed.pages == []
    assert parsed.pages_total is None
